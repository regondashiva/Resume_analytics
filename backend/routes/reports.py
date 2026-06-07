from fastapi import APIRouter, Depends, HTTPException, status, Response
from sqlalchemy.orm import Session
from models.models import User, MatchingResult, Job, Candidate
from models.schemas import ReportGenerateRequest, ExportCandidatesRequest
from database.connection import get_db
from authentication.jwt_handler import get_current_user
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib import colors
import io
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

router = APIRouter(prefix="/report", tags=["reports"])


@router.post("/generate")
async def generate_report(
    payload: ReportGenerateRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Generate recruitment report"""
    job_id = payload.job_id
    if job_id == "":
        job_id = None

    if payload.format == "pdf":
        pdf_data = await generate_pdf_report(db, job_id)
        return Response(
            content=pdf_data,
            media_type="application/pdf",
            headers={"Content-Disposition": "attachment; filename=recruitment-report.pdf"}
        )
    elif payload.format == "excel":
        excel_data = await generate_excel_report(db, job_id)
        return Response(
            content=excel_data,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=recruitment-report.xlsx"}
        )
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Invalid format"
        )


async def generate_pdf_report(db: Session, job_id: str):
    """Generate PDF report"""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    elements = []

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=24,
        textColor=colors.HexColor('#3B82F6'),
        spaceAfter=30,
        alignment=1,  # Center
     )

    # Title
    elements.append(Paragraph("Recruitment Analytics Report", title_style))
    elements.append(Spacer(1, 0.3 * inch))

    # Get matching results
    if job_id:
        matches = db.query(MatchingResult).filter(
            MatchingResult.job_id == job_id
        ).order_by(MatchingResult.match_score.desc()).all()
    else:
        matches = db.query(MatchingResult).order_by(
            MatchingResult.match_score.desc()
        ).limit(50).all()

    # Create table data
    table_data = [["Rank", "Candidate", "Match Score", "Recommendation"]]
    for idx, match in enumerate(matches[:20], 1):
        candidate = db.query(Candidate).filter(
            Candidate.id == match.candidate_id
        ).first()
        rec_str = (match.recommendation or "N/A").upper()
        score_val = match.match_score or 0.0
        table_data.append([
            str(idx),
            candidate.name if candidate else "N/A",
            f"{score_val:.2f}%",
            rec_str,
        ])

    # Create table
    if len(table_data) > 1:
        table = Table(table_data)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#3B82F6')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ]))
        elements.append(table)

    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()


async def generate_excel_report(db: Session, job_id: str):
    """Generate Excel report"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Recruitment Report"

    # Add headers
    headers = ["Rank", "Candidate Name", "Email", "Skills", "Experience", "Match Score", "Recommendation"]
    ws.append(headers)

    # Style headers
    header_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Get data
    if job_id:
        matches = db.query(MatchingResult).filter(
            MatchingResult.job_id == job_id
        ).order_by(MatchingResult.match_score.desc()).all()
    else:
        matches = db.query(MatchingResult).order_by(
            MatchingResult.match_score.desc()
        ).limit(100).all()

    # Add rows
    for idx, match in enumerate(matches, 1):
        candidate = db.query(Candidate).filter(
            Candidate.id == match.candidate_id
        ).first()

        if candidate:
            skills_list = candidate.skills or []
            skills_str = ", ".join(skills_list) if isinstance(skills_list, list) else str(skills_list)
            exp_val = candidate.experience_years or 0.0
            score_val = match.match_score or 0.0
            rec_str = (match.recommendation or "N/A").upper()
            
            ws.append([
                idx,
                candidate.name,
                candidate.email,
                skills_str,
                f"{exp_val:.1f}",
                f"{score_val:.2f}%",
                rec_str,
            ])

    # Adjust column widths
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 30
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 15
    ws.column_dimensions['G'].width = 15

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


@router.post("/export-candidates")
async def export_candidates(
    payload: ExportCandidatesRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Export all candidates"""
    candidates = db.query(Candidate).all()

    if payload.format == "excel":
        wb = Workbook()
        ws = wb.active
        ws.title = "Candidates"

        # Add headers
        headers = ["ID", "Name", "Email", "Phone", "Skills", "Experience (Years)", "Score", "Education"]
        ws.append(headers)

        # Style headers
        header_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # Add rows
        for candidate in candidates:
            skills_list = candidate.skills or []
            skills_str = ", ".join(skills_list) if isinstance(skills_list, list) else str(skills_list)
            exp_val = candidate.experience_years or 0.0
            score_val = candidate.score or 0.0
            
            ws.append([
                candidate.id,
                candidate.name,
                candidate.email,
                candidate.phone or "",
                skills_str,
                f"{exp_val:.1f}",
                f"{score_val:.2f}",
                candidate.education or "",
            ])

        # Adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[column_letter].width = max_length + 2

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        return Response(
            content=buffer.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=candidates-export.xlsx"}
        )

    elif payload.format == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["ID", "Name", "Email", "Phone", "Skills", "Experience (Years)", "Score", "Education"])
        
        for candidate in candidates:
            skills_list = candidate.skills or []
            skills_str = ", ".join(skills_list) if isinstance(skills_list, list) else str(skills_list)
            exp_val = candidate.experience_years or 0.0
            score_val = candidate.score or 0.0
            
            writer.writerow([
                candidate.id,
                candidate.name,
                candidate.email,
                candidate.phone or "",
                skills_str,
                f"{exp_val:.1f}",
                f"{score_val:.2f}",
                candidate.education or "",
            ])
            
        csv_data = output.getvalue().encode('utf-8')
        return Response(
            content=csv_data,
            media_type="text/csv",
            headers={"Content-Disposition": "attachment; filename=candidates-export.csv"}
        )

    raise HTTPException(
        status_code=status.HTTP_400_BAD_REQUEST,
        detail="Invalid format. Supported: excel, csv"
    )

